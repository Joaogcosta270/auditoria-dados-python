import pandas as pd
from sqlalchemy import create_engine, text
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

conexao_url = r"mssql+pyodbc://localhost\SQLEXPRESS/auditoria?driver=ODBC+Driver+17+for+SQL+Server&trusted_connection=yes"
engine = create_engine(conexao_url)

def rodar_auditoria():
    print("Iniciando auditoria cruzada...")

    try:
        query_sql = "SELECT id_venda, cliente, valor_original, status_sistema FROM vendas"
        df_sql = pd.read_sql_query(query_sql, engine)
        print("-> Dados do SQL Server carregados com sucesso!")
    except Exception as e:
        print(f"Erro ao ler dados do SQL Server: {e}")
        return

    try:
        df_excel = pd.read_excel('planilha_financeiro.xlsx')
        print("-> Planilha do Financeiro carregada com sucesso!")
    except FileNotFoundError:
        print("Erro: O arquivo 'planilha_financeiro.xlsx' não foi encontrado na pasta.")
        return

    relatorio_erros = []
    ids_corretos_para_atualizar = []

    duplicados = df_excel[df_excel.duplicated(subset=['id_venda'], keep=False)]
    ids_duplicados = set(duplicados['id_venda'])

    for idx, row in duplicados.iterrows():
        relatorio_erros.append({
            'ID Venda': row['id_venda'],
            'Cliente': row['cliente'],
            'Tipo de Erro': 'Linha Duplicada na Planilha',
            'Detalhe': 'O ID desta venda aparece repetido no Excel enviado.'
        })

    df_excel_clean = df_excel.drop_duplicates(subset=['id_venda'], keep='first')
    df_cruzado = pd.merge(df_sql, df_excel_clean, on='id_venda', suffixes=('_sql', '_excel'), how='left')

    for idx, row in df_cruzado.iterrows():
        id_venda = row['id_venda']

        if id_venda in ids_duplicados:
            continue

        if pd.isna(row['valor_pago']):
            relatorio_erros.append({
                'ID Venda': id_venda,
                'Cliente': row['cliente_sql'],
                'Tipo de Erro': 'Venda Ausente',
                'Detalhe': 'Consta no SQL Server, mas sumiu da planilha do financeiro.'
            })
            continue

        if row['cliente_sql'] != row['cliente_excel']:
            relatorio_erros.append({
                'ID Venda': id_venda,
                'Cliente': row['cliente_excel'],
                'Tipo de Erro': 'Nome Divergente',
                'Detalhe': f"No banco: '{row['cliente_sql']}'. No Excel: '{row['cliente_excel']}'."
            })

        if float(row['valor_original']) != float(row['valor_pago']):
            relatorio_erros.append({
                'ID Venda': id_venda,
                'Cliente': row['cliente_excel'],
                'Tipo de Erro': 'Valor Divergente',
                'Detalhe': f"Esperado no banco: R${row['valor_original']}. Pago no Excel: R${row['valor_pago']}."
            })
        else:
            ids_corretos_para_atualizar.append(int(id_venda))

    if ids_corretos_para_atualizar:
        try:
            with engine.begin() as conn_update:
                for id_v in ids_corretos_para_atualizar:
                    conn_update.execute(
                        text("UPDATE vendas SET status_sistema = 'Confirmado e Pago' WHERE id_venda = :id"),
                        {"id": id_v}
                    )
            print(f"-> SQL Server atualizado com sucesso! {len(ids_corretos_para_atualizar)} registros validados.")
        except Exception as e:
            print(f"Erro ao atualizar o SQL Server: {e}")
    else:
        print("-> Nenhuma venda correta encontrada para atualizar no sistema.")

    if relatorio_erros:
        df_erros = pd.DataFrame(relatorio_erros)
        nome_relatorio = 'Relatorio_Auditoria_Erros.xlsx'
        df_erros.to_excel(nome_relatorio, index=False)

        wb = load_workbook(nome_relatorio)
        ws = wb.active
        vermelho_alerta = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.fill = vermelho_alerta

        wb.save(nome_relatorio)
        print(f"-> Relatório de erros gerado com sucesso: '{nome_relatorio}'")
    else:
        print("-> Nenhum erro encontrado nos arquivos.")

if __name__ == '__main__':
    rodar_auditoria()